"""
Handles reading, writing, and placeholder replacement in Word documents.
Requires: python-docx (install via pip if needed)
"""

from docx import Document

class DocxHandler:
    def __init__(self, template_path):
        self.template_path = template_path
        self.document = Document(template_path)

    def replace_placeholders(self, replacements):
        """
        replacements: dict mapping placeholder syntax to replacement value
        """
        for para in self.document.paragraphs:
            for placeholder, value in replacements.items():
                if placeholder in para.text:
                    para.text = para.text.replace(placeholder, value)

    def save(self, output_path):
        self.document.save(output_path)